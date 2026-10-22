"""Converte a folha de cálculo antiga (Social Blade à mão) em public/data/seed.json.

Uso:
    python3 scripts/seed_from_xlsx.py caminho/para/folha.xlsx

O seed só é usado na primeira vez que o site arranca (base de dados vazia).
Depois disso, tudo é editado na área de edição do site.
"""
import json
import re
import sys
import unicodedata
from datetime import datetime, date, timezone
from pathlib import Path

import openpyxl

MONTHS = [
    "2024-12", "2025-01", "2025-02", "2025-03", "2025-04", "2025-05", "2025-06",
    "2025-07", "2025-08", "2025-09", "2025-10", "2025-11", "2025-12", "2026-01",
    "2026-02", "2026-03",
]

# Variantes sem acento que existem na folha -> nome final da categoria
CATEGORY_FIX = {
    "Politica": "Política",
    "Comedia": "Comédia",
    "Musica": "Música",
    "Historia": "História",
    "Noticias": "Notícias",
    "Danca": "Dança",
    "Fisica": "Física",
    "Educacao": "Educação",
}


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[\s  ]", "", str(v))
    try:
        return float(s)
    except ValueError:
        return None


def slug(name):
    s = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-zA-Z0-9]+", "-", s).strip("-").lower()
    return s or "p"


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    dados = wb["Dados"]
    yt = wb["Youtube"]

    info = {}
    for r in dados.iter_rows(min_row=2, values_only=True):
        if not r[2]:
            continue
        info[r[2]] = r

    people = []
    used = set()
    for r in yt.iter_rows(min_row=3, values_only=True):
        name = r[1]
        if not name:
            continue
        d = info.get(name)
        pid = slug(name)
        base, n = pid, 2
        while pid in used:
            pid = f"{base}-{n}"
            n += 1
        used.add(pid)

        monthly = {}
        for i, m in enumerate(MONTHS):
            v = num(r[3 + i])
            if v is not None and v >= 0:
                monthly[m] = {"v": int(round(v)), "src": "folha"}

        cat = (d[31] if d else r[21]) or ""
        cat = CATEGORY_FIX.get(cat, cat)
        legacy = {}
        if d:
            if isinstance(d[32], (int, float)):
                legacy["pontuacao"] = int(d[32])
            if num(d[33]):
                legacy["totalViews"] = int(num(d[33]))
            if num(d[34]):
                legacy["subs"] = int(num(d[34]))
            if isinstance(d[35], (datetime, date)):
                legacy["createdAt"] = d[35].strftime("%Y-%m-%d")

        people.append({
            "id": pid,
            "name": name,
            "category": cat,
            "monthly": monthly,
            "legacy": legacy,
        })

    out = {
        "version": 1,
        "source": "Folha de cálculo (Social Blade, recolha manual)",
        "generatedAt": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "people": people,
    }
    dest = Path(__file__).resolve().parent.parent / "public" / "data" / "seed.json"
    dest.write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    cats = sorted({p["category"] for p in people})
    print(f"{len(people)} pessoas, {len(cats)} categorias -> {dest}")


if __name__ == "__main__":
    main(sys.argv[1])
